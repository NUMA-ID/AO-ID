#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Parseur d'export Dell Solutions Configurator (.xlsx) ou de devis distributeur
(ex. TD SYNNEX) -> fiche de specs JSON ONE ID.

Deux formats reconnus (auto-détection par en-tête) :

1. Dell Solutions Configurator : tableau hierarchique unique :
   - lignes d'en-tete solution (ID, Nom, Categorie, Prix...)
   - ligne de colonnes : Nom du groupe | ID groupe | Nom du produit | Quantite |
     Prix unitaire | Prix etendu | Nom du module | ID option | Nom option | Prix | SKU | Qte
   - produits (colonne "Nom du produit" renseignee)
   - composants/options de chaque produit (colonne "Nom du module" renseignee)

2. Devis distributeur (TD SYNNEX) : tableau plus simple :
   - ligne de colonnes : Référence Constructeur | Réf. TD SYNNEX | Description |
     Prix public unit. HT EUR | Total List Price EUR | Remise | Prix d'achat
     unit. HT EUR | Qté | Prix Total HT EUR | ...
   - lignes produit : "Référence Constructeur" + "Description" renseignés
   - lignes composant/option : une seule cellule "Référence Constructeur"
     au format "<N>X <libellé>" (ex. "2X 480GB SSD SATA...")

Sortie : JSON compatible avec generer_doc.py, categories :
  serveurs / stockages / switches / sauvegardes / logiciels_services

Usage : python3 parser_dell_excel.py entree.xlsx sortie.json
"""
import sys, json, re, datetime, unicodedata
import openpyxl

# --- Correctif lecture : certains exports (dont les devis TD SYNNEX observés en
# production) contiennent un attribut de style <font family="…"> supérieur à 14.
# openpyxl (3.1.5) applique une validation trop stricte (max=14, voir
# openpyxl/styles/fonts.py, descripteur Font.family) alors que la norme OOXML
# n'impose pas cette limite en pratique. Sans ce correctif, l'ouverture du
# classeur échoue avec "ValueError: Max value is 14" dès load_workbook().
# Voir docs/BLOCAGES.md, incident du 2026-08-24 (PROD, import Excel HTTP 500).
try:
    import openpyxl.styles.fonts as _openpyxl_fonts
    _openpyxl_fonts.Font.family.max = 999
except Exception:
    pass


class HeaderNotFoundError(ValueError):
    """Levée quand aucun format d'en-tête connu (Dell ou TD SYNNEX) n'est trouvé."""


def deacc(s):
    return "".join(c for c in unicodedata.normalize("NFD", s)
                   if unicodedata.category(c) != "Mn")

# Classification par mots-cles (ordre = priorite)
RULES = [
    ("logiciels_services", ["appsync", "required software", "data manager", "cloudiq", "veeam", "commvault"]),
    ("sauvegardes", ["powerprotect", "data domain", "datadomain", " dd3", " dd6", " dd9",
                     "ddos", "avamar", "networker", "cyber recovery", "ppdm",
                     "rapidrecovery", "idpa"]),
    ("switches",    ["powerswitch", "networking", "connectrix", "brocade", "switch",
                     " s5", " s4", " s3", " n3", " z9", " z10", "mx9116", "mx5108",
                     "ds-", "sfp"]),
    ("stockages",   ["powerstore", "unity", "powervault", "me5", "me4", "powerscale",
                     "isilon", "powermax", "vnx", "scv", "ecs", "objectscale",
                     "storage", "stockage", "baie"]),
    ("serveurs",    ["poweredge", "pe r", "pe t", " r6", " r7", " r8", " r9", " t1",
                     " t3", " t5", "mx740", "mx760", "mx840", "fc640", "vxrail",
                     "serveur", "server"]),
]
CAT_LABEL = {
    "serveurs": "Serveurs", "stockages": "Stockage", "switches": "Switches",
    "sauvegardes": "Baies de sauvegarde", "logiciels_services": "Logiciels & Services",
}
HARDWARE = ("serveurs", "stockages", "switches", "sauvegardes")

SPEC_HINTS = [
    ("disques", ["drive", "disque", "hard drive", "ssd", "hdd"]),
    ("memoire", ["memory", "dimm", "memoire", "ram"]),
    ("alimentation", ["power supply", "alimentation", "bloc d"]),
    ("reseau", ["mezz", "nic", "carte", "ocp", "ethernet", "gbe", " port", "sfp"]),
    ("processeur", ["processor", "processeur", "cpu", "xeon", "epyc"]),
    ("os", ["operating", "systeme", "ddos", "logiciel de base", "base logiciel"]),
    ("capacite", ["capacity", "capacite"]),
    ("support", ["prosupport", "support", "garantie", "service", "warranty"]),
]

# Ligne composant TD SYNNEX : "<N>X <libellé>" (ex. "24X Informational Purposes Only").
COMP_QTE_RE = re.compile(r"^(\d+)\s*[xX]\s+(.+)$")


def norm(v):
    if v is None:
        return ""
    s = str(v).replace("\n", " ").strip()
    return re.sub(r"\s+", " ", s)


def classify_text(text):
    t = " " + (text or "").lower() + " "
    for cat, kws in RULES:
        if any(k in t for k in kws):
            return cat
    return None


def classify(p):
    if isinstance(p, str):
        return classify_text(p) or "logiciels_services"
    return classify_text(p["produit"]) or classify_text(ptext(p)) or "logiciels_services"


def ptext(p):
    return p["produit"] + " " + " ".join(c["module"] for c in p["composants"])


def find_header(ws):
    """Détecte l'en-tête format Dell Solutions Configurator."""
    for r in range(1, min(ws.max_row, 30) + 1):
        rowvals = [deacc(norm(ws.cell(r, c).value).lower()) for c in range(1, ws.max_column + 1)]
        if "nom du produit" in rowvals and "nom du module" in rowvals:
            idx = {}
            for c in range(1, ws.max_column + 1):
                v = deacc(norm(ws.cell(r, c).value).lower())
                if "nom du groupe" in v:
                    idx["groupe"] = c
                elif "nom du produit" in v:
                    idx["produit"] = c
                elif v.startswith("quantite du produit") or v == "quantite":
                    idx["qte_prod"] = c
                elif "prix unitaire" in v:
                    idx["pu"] = c
                elif "prix etendu" in v:
                    idx["pe"] = c
                elif "nom du module" in v:
                    idx["module"] = c
                elif "nom de l" in v and "option" in v:
                    idx["option"] = c
                elif v == "sku":
                    idx["sku"] = c
                elif v == "qte":
                    idx["qte_mod"] = c
                elif "prix catalogue" in v and "option" in v:
                    idx["prix_mod"] = c
            return r, idx
    raise HeaderNotFoundError("En-tete Dell introuvable (colonnes 'Nom du produit'/'Nom du module').")


def find_header_tdsynnex(ws):
    """Détecte l'en-tête format devis distributeur (TD SYNNEX)."""
    for r in range(1, min(ws.max_row, 40) + 1):
        rowvals = [deacc(norm(ws.cell(r, c).value).lower()) for c in range(1, ws.max_column + 1)]
        if "reference constructeur" in rowvals and "description" in rowvals and "qte" in rowvals:
            idx = {}
            for c in range(1, ws.max_column + 1):
                v = deacc(norm(ws.cell(r, c).value).lower())
                if v == "reference constructeur":
                    idx["ref"] = c
                elif v == "description":
                    idx["description"] = c
                elif "prix public" in v:
                    idx["prix_public"] = c
                elif "achat" in v:
                    idx["prix_achat"] = c
                elif v == "qte":
                    idx["qte"] = c
                elif "prix total" in v and "ht" in v:
                    idx["prix_total"] = c
            return r, idx
    raise HeaderNotFoundError(
        "En-tete TD SYNNEX introuvable (colonnes 'Reference Constructeur'/'Description'/'Qte').")


def fmt_qte(val, qte):
    if qte and qte not in ("", "1"):
        return val + " [qte " + str(qte) + "]"
    return val


def build_item(p, cat, qte):
    specs = {}
    for comp in p["composants"]:
        label = (comp["module"] + " " + comp["description"]).lower()
        for key, kws in SPEC_HINTS:
            if any(k in label for k in kws):
                val = comp["description"] or comp["module"]
                val = fmt_qte(val, comp.get("qte"))
                specs.setdefault(key, []).append(val)
                break
    grp = (p["groupe"] or "").strip()
    if re.match(r"(?i)^group\s*\d+$", grp):
        grp = ""
    return {
        "_source": "datasheet",
        "role": grp,
        "modele": p["produit"],
        "qte": qte,
        "prix_unitaire": p["prix_unitaire"],
        "prix_etendu": p["prix_etendu"],
        "specs_extraites": dict((k, " ; ".join(v)) for k, v in specs.items()),
        "composants": p["composants"],
    }


def parse_dell(path, ws, hdr, idx):
    sol = {}
    for r in range(1, hdr):
        a = norm(ws.cell(r, 1).value)
        if ":" in a:
            k, _, v = a.partition(":")
            sol[k.strip().lower()] = v.strip()

    def cell(r, key):
        return norm(ws.cell(r, idx[key]).value) if key in idx else ""

    groupes = []
    current_grp = None
    current_prod = None
    for r in range(hdr + 1, ws.max_row + 1):
        grp = cell(r, "groupe")
        prod = cell(r, "produit")
        mod = cell(r, "module")
        if grp:
            current_grp = {"label": grp, "qte": cell(r, "qte_prod") or "1",
                           "prix_etendu": cell(r, "pe"), "produits": []}
            groupes.append(current_grp)
        if prod:
            if current_grp is None:
                current_grp = {"label": "", "qte": "1", "prix_etendu": "", "produits": []}
                groupes.append(current_grp)
            current_prod = {
                "produit": prod, "groupe": current_grp["label"],
                "qte": cell(r, "qte_prod") or "1",
                "prix_unitaire": cell(r, "pu"), "prix_etendu": cell(r, "pe"),
                "composants": [],
            }
            current_grp["produits"].append(current_prod)
        elif mod and current_prod is not None:
            current_prod["composants"].append({
                "module": mod, "description": cell(r, "option"),
                "sku": cell(r, "sku"), "qte": cell(r, "qte_mod"), "prix": cell(r, "prix_mod"),
            })

    buckets = dict((k, []) for k in CAT_LABEL)
    for g in groupes:
        prods = [p for p in g["produits"] if p["composants"]]
        primary = next((p for p in prods if classify(p) in HARDWARE), None)
        if primary is None and prods:
            primary = prods[0]
        for p in prods:
            cat = classify(p)
            qte = g["qte"] if (p is primary and g["qte"]) else p["qte"]
            buckets[cat].append(build_item(p, cat, qte))

    return {
        "_type": "fiche_specs_oneid", "_version": 2,
        "_import": {"source": "Dell Solutions Configurator", "fichier": path,
                    "le": datetime.date.today().isoformat()},
        "affaire": {
            "client": "", "projet": sol.get("nom de la solution", ""),
            "reference": sol.get("id de la solution", ""),
            "version": "1.0", "auteur": "", "date": datetime.date.today().isoformat(),
            "notes": sol.get("infos sur la solution", ""),
            "prix_solution": sol.get("prix catalogue de la solution", ""),
        },
        "serveurs": buckets["serveurs"],
        "stockages": buckets["stockages"],
        "switches": buckets["switches"],
        "sauvegardes": buckets["sauvegardes"],
        "logiciels_services": buckets["logiciels_services"],
    }


def parse_tdsynnex(path, ws, hdr, idx):
    def cell(r, key):
        return norm(ws.cell(r, idx[key]).value) if key in idx else ""

    def find_meta(label_norm):
        """Cherche un libellé (ex. 'client final') avant l'en-tête et renvoie la
        première cellule non vide qui le suit sur la même ligne."""
        for r in range(1, hdr):
            for c in range(1, ws.max_column + 1):
                v = deacc(norm(ws.cell(r, c).value).lower()).rstrip(":").strip()
                if v == label_norm:
                    for c2 in range(c + 1, ws.max_column + 1):
                        vv = norm(ws.cell(r, c2).value)
                        if vv:
                            return vv
        return ""

    client = find_meta("client final")
    devis = find_meta("numero de devis")

    produits = []
    current = None
    for r in range(hdr + 1, ws.max_row + 1):
        desc = cell(r, "description")
        ref1 = norm(ws.cell(r, 1).value)  # colonne A : réf. produit OU ligne composant "NX ..."
        if desc:
            current = {
                "produit": desc, "groupe": "",
                "qte": cell(r, "qte") or "1",
                "prix_unitaire": cell(r, "prix_achat") or cell(r, "prix_public"),
                "prix_etendu": cell(r, "prix_total"),
                "composants": [],
            }
            produits.append(current)
        elif ref1 and current is not None:
            m = COMP_QTE_RE.match(ref1)
            if m:
                current["composants"].append({
                    "module": "", "description": m.group(2),
                    "sku": "", "qte": m.group(1), "prix": "",
                })

    buckets = dict((k, []) for k in CAT_LABEL)
    for p in produits:
        if not p["composants"]:
            continue
        cat = classify(p)
        buckets[cat].append(build_item(p, cat, p["qte"]))

    return {
        "_type": "fiche_specs_oneid", "_version": 2,
        "_import": {"source": "Devis TD SYNNEX", "fichier": path,
                    "le": datetime.date.today().isoformat()},
        "affaire": {
            "client": client, "projet": devis,
            "reference": devis, "version": "1.0", "auteur": "",
            "date": datetime.date.today().isoformat(),
            "notes": "", "prix_solution": "",
        },
        "serveurs": buckets["serveurs"],
        "stockages": buckets["stockages"],
        "switches": buckets["switches"],
        "sauvegardes": buckets["sauvegardes"],
        "logiciels_services": buckets["logiciels_services"],
    }


def parse(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    try:
        hdr, idx = find_header(ws)
        return parse_dell(path, ws, hdr, idx)
    except HeaderNotFoundError:
        pass
    try:
        hdr, idx = find_header_tdsynnex(ws)
        return parse_tdsynnex(path, ws, hdr, idx)
    except HeaderNotFoundError:
        raise HeaderNotFoundError(
            "Format de fichier non reconnu : ni export Dell Solutions Configurator "
            "(colonnes 'Nom du produit'/'Nom du module'), ni devis TD SYNNEX "
            "(colonnes 'Référence Constructeur'/'Description'/'Qté').")


def main():
    if len(sys.argv) < 3:
        print("Usage: python3 parser_dell_excel.py entree.xlsx sortie.json")
        sys.exit(1)
    fiche = parse(sys.argv[1])
    with open(sys.argv[2], "w", encoding="utf-8") as f:
        json.dump(fiche, f, ensure_ascii=False, indent=2)
    for cat in CAT_LABEL:
        items = fiche.get(cat, [])
        if items:
            print(CAT_LABEL[cat] + " : " + str(len(items)))
            for it in items:
                print("   - " + it["modele"] + " (x" + str(it["qte"]) + ") - "
                      + str(len(it["composants"])) + " composants")
    print("->", sys.argv[2])


if __name__ == "__main__":
    main()

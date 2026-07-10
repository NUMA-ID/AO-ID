#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Generateur de document technique ONE ID (appel d'offre).
Sortie .docx au modele ONE ID : intro engageante, synthese de la solution,
chapitres par famille, argumentaires technico-commerciaux sources.
Le contenu PERSONNALISE (specifique au dossier) est mis en ROUGE pour relecture.
Usage : python3 generer_doc.py fiche.json sortie.docx [Documentation_Constructeur]"""
import sys, json, zipfile, shutil, os, datetime, re, glob, base64, io
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))

def find_base():
    cands = [os.environ.get("BASE_DOCX"),
             os.path.join(HERE, "TEMPLATE_BASE_ONEID.docx"),
             os.path.join(HERE, "..", "assets", "TEMPLATE_BASE_ONEID.docx")]
    for c in cands:
        if c and os.path.exists(c):
            return c
    return os.path.join(HERE, "TEMPLATE_BASE_ONEID.docx")

BASE = find_base()
H1, H2, H3 = "Heading 1", "Heading 2", "Heading 3"
BODY = "Paragraphe_Exakis"
TABLE_STYLE = "Light List Accent 6"
RED = RGBColor(0xC0, 0x00, 0x00)
ONEID = RGBColor(0x0a, 0x3d, 0x62)
FONT_NAME = "Assistant"          # police du corps de texte
FONT_BODY_PT = 11                # taille par défaut du texte
ARG_KEYWORDS = ["PowerStore", "PowerProtect", "Data Domain", "DataDomain", "Veeam", "AirGap",
                "Air Gap", "NIS2", "Zero-Trust", "Enterprise Plus", "Metro", "CDP",
                "déduplication", "immuables", "immuable", "FortiGate", "PowerEdge", "ProSupport"]
KW_RE = re.compile("(" + "|".join(re.escape(k) for k in ARG_KEYWORDS) + ")", re.IGNORECASE)
HDR_HEX = "0a3d62"      # couleur d'en-têtes/cases (modifiable par le thème Design)
DESIGN_ON = False       # mode Design actif → titres colorés
_H1_SEEN = False        # saut de page avant chaque chapitre principal (Titre 1), sauf le 1er


def apply_theme(theme):
    """Applique un thème Design (couleur primaire) ; sans thème, rendu standard inchangé."""
    global ONEID, HDR_HEX, DESIGN_ON
    if not theme:
        return
    prim = str(theme.get("primary") or "").lstrip("#")
    if len(prim) == 6:
        try:
            ONEID = RGBColor.from_string(prim.upper()); HDR_HEX = prim.lower(); DESIGN_ON = True
        except Exception:
            pass
MOIS = ["janvier","fevrier","mars","avril","mai","juin","juillet","aout",
        "septembre","octobre","novembre","decembre"]

CATS = [
    ("serveurs", "Serveurs", "Les serveurs assurent la puissance de calcul de votre infrastructure.", [
        ("Rôle / désignation","role"),("Quantité","qte"),("Modèle","modele"),
        ("Facteur de forme","format"),("CPU (famille)","cpu_fam"),("Nombre de sockets","cpu_sockets"),
        ("CPU souhaité","cpu_detail"),("Mémoire RAM","ram"),("Type de disques","disq_type"),
        ("Configuration disques","disq_detail"),("RAID","raid"),("Réseau","reseau"),
        ("Alimentation","alim"),("Système d'exploitation","os"),("Garantie / support","garantie"),
        ("Remarques","remarques")]),
    ("stockages", "Stockage", "Le stockage héberge et protège vos données avec performance et évolutivité.", [
        ("Désignation","role"),("Quantité","qte"),("Modèle","modele"),("Mode","mode"),
        ("Type de disque","disq_type"),("Protocoles","protocoles"),("Taille RAW","raw"),
        ("Taille utile requise","utile"),("Compression / Déduplication","dedup"),
        ("Protection / RAID","raid"),("Contrôleurs","controleurs"),("Connectique / ports","ports"),
        ("Évolutivité","eval"),("Garantie / support","garantie"),("Remarques","remarques")]),
    ("switches", "Switches", "Le réseau interconnecte l'ensemble des composants de la solution.", [
        ("Désignation","role"),("Quantité","qte"),("Modèle","modele"),("Nombre / type de ports","ports"),
        ("Débit","debit"),("PoE","poe"),("Empilage / stacking","stacking"),("Niveau (L2/L3)","l2l3"),
        ("Garantie / support","garantie"),("Remarques","remarques")]),
    ("sauvegardes", "Baies de sauvegarde", "La sauvegarde garantit la reprise et la résilience de vos données.", [
        ("Désignation","role"),("Quantité","qte"),("Modèle","modele"),("Type","type"),
        ("Capacité brute","raw"),("Capacité utile","utile"),("Déduplication / compression","dedup"),
        ("Rétention","retention"),("Protocoles","protocoles"),("Garantie / support","garantie"),
        ("Remarques","remarques")]),
    ("logiciels_services", "Logiciels & Services", "Les logiciels et services complètent et sécurisent la solution.", [
        ("Désignation","role"),("Quantité","qte"),("Produit","modele"),("Remarques","remarques")]),
]
SPEC_LABEL = {"processeur":"Processeur","memoire":"Mémoire","disques":"Disques","capacite":"Capacité",
    "reseau":"Réseau / connectique","alimentation":"Alimentation","os":"Système / logiciel","support":"Support / garantie"}
SRC_LABEL = {"internet":"Internet (site constructeur)","datasheet":"Caractéristiques constructeur (devis Dell)","both":"Comparaison Internet + Datasheet"}

ARGS, SHOWN_ARGS, AFF, DOC_BASE = [], set(), {}, ""


def sstyle(doc, p, s):
    try: p.style = doc.styles[s]
    except KeyError: pass
    return p

def add_para(doc, text="", style=BODY, red=False, bold=False, italic=False, size=None, color=None):
    p = doc.add_paragraph(); sstyle(doc, p, style)
    if text:
        r = p.add_run(text)
        if color: r.font.color.rgb = RGBColor.from_string(color)
        elif red: r.font.color.rgb = RED
        if bold: r.bold = True
        if italic: r.italic = True
        _afont(r, size)
    return p

def add_heading(doc, text, level, qte=None):
    global _H1_SEEN
    p = doc.add_paragraph(); sstyle(doc, p, {1:H1,2:H2,3:H3}[level])
    if level == 1:
        if _H1_SEEN:
            p.paragraph_format.page_break_before = True
        _H1_SEEN = True
    r0 = p.add_run(text)
    if DESIGN_ON:
        r0.font.color.rgb = ONEID
    if qte not in (None, "", "1"):
        r = p.add_run("  (x" + str(qte) + ")"); r.font.color.rgb = RED
    return p

def _fill_runs(p, text, ctx):
    """Ajoute le texte au paragraphe : accroche \"Label : \" en gras, {clé} en ROUGE."""
    rest = text
    m = re.match(r"^([A-Z\u00C0-\u00DD][^:{}.]{2,60}?)\s:\s(.*)$", text, re.S)
    if m:
        lead = p.add_run(m.group(1) + " : "); lead.bold = True; _afont(lead)
        rest = m.group(2)
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", rest)):
        if part == "":
            continue
        _emit(p, part, ctx, force_bold=(i % 2 == 1))


def _afont(run, size=None):
    run.font.name = FONT_NAME
    run.font.size = Pt(size or FONT_BODY_PT)
    return run


def _emit(p, text, ctx, force_bold=False):
    for seg in re.split(r"(\{\w+\})", text):
        if seg == "":
            continue
        mm = re.fullmatch(r"\{(\w+)\}", seg)
        if mm:
            r = p.add_run(str(ctx.get(mm.group(1), seg))); r.font.color.rgb = RED
            if force_bold: r.bold = True
            _afont(r)
        elif force_bold:
            _afont(p.add_run(seg)).bold = True
        else:
            for part in KW_RE.split(seg):
                if not part:
                    continue
                r = p.add_run(part)
                if KW_RE.fullmatch(part): r.bold = True
                _afont(r)

def add_subst(doc, template, ctx, style=BODY):
    p = doc.add_paragraph(); sstyle(doc, p, style); _fill_runs(p, template, ctx); return p

def style_table(doc, t):
    for st in (TABLE_STYLE, "Table Grid"):
        try: t.style = doc.styles[st]; return
        except KeyError: continue

def set_cell(cell, text, red=False, bold=False, size=None):
    cell.text = ""
    r = cell.paragraphs[0].add_run("" if text is None else str(text))
    if red: r.font.color.rgb = RED
    if bold: r.bold = True
    _afont(r)
    if size: r.font.size = Pt(size)

def kv_table(doc, rows, headers=("Caractéristique","Valeur"), red_values=False, widths=None):
    rows = [(k, v) for k, v in rows if v not in (None,"","À préciser")]
    if not rows: return
    t = doc.add_table(rows=1, cols=2); style_table(doc, t)
    set_cell(t.rows[0].cells[0], headers[0], bold=True); set_cell(t.rows[0].cells[1], headers[1], bold=True)
    for k, v in rows:
        c = t.add_row().cells
        set_cell(c[0], k); set_cell(c[1], v, red=red_values)
    _set_widths(t, widths or [Cm(3.4), Cm(13.0)])
    return t

def nomenclature_table(doc, comps):
    comps = [c for c in comps if c.get("description") or c.get("module")]
    if not comps: return
    t = doc.add_table(rows=1, cols=4); style_table(doc, t)
    for i, h in enumerate(("Module","Description","Référence (SKU)","Qté")):
        set_cell(t.rows[0].cells[i], h, bold=True, size=8)
    for c in comps:
        r = t.add_row().cells
        set_cell(r[0], c.get("module",""), size=8); set_cell(r[1], c.get("description",""), size=8)
        set_cell(r[2], c.get("sku",""), size=8); set_cell(r[3], c.get("qte",""), size=8)
    _set_widths(t, [Cm(3.5), Cm(10.0), Cm(2.0), Cm(1.0)])

def insert_images(doc, item, arg=None):
    placed = False
    # 1) images fournies dans la fiche (base64)
    for im in (item.get("images") or []):
        data = im.get("data") if isinstance(im, dict) else im
        if not data: continue
        if "," in data and "base64" in data[:40]: data = data.split(",", 1)[1]
        try:
            doc.add_picture(io.BytesIO(base64.b64decode(data)), width=Cm(13))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER; placed = True
        except Exception: pass
    # 2) sinon, image(s) par defaut du produit (fichiers du dossier de l'argumentaire)
    if not placed and arg:
        for fn in arg.get("images_defaut", []):
            for fp in (os.path.join(arg.get("_dir", ""), fn), os.path.join(os.path.dirname(BASE), fn)):
                if os.path.exists(fp):
                    try:
                        doc.add_picture(fp, width=Cm(13))
                        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER; placed = True
                        break
                    except Exception: pass
    # 3) sinon, emplacement reserve
    if not placed:
        ph = add_para(doc, "[ Image produit à insérer ]", red=True, italic=True)
        ph.alignment = WD_ALIGN_PARAGRAPH.CENTER

def load_argumentaires():
    base = DOC_BASE or os.path.join(os.getcwd(), "Documentation_Constructeur")
    out = []
    for path in glob.glob(os.path.join(base, "*", "argumentaire.json")):
        try:
            a = json.load(open(path, encoding="utf-8")); a["_dir"] = os.path.dirname(path); out.append(a)
        except Exception: pass
    return out

def match_arg(args, modele):
    m = (modele or "").lower()
    for a in args:
        if any(str(k).lower() in m for k in a.get("modele_match", [])): return a
    return None

def ctx_item(item):
    client = (AFF.get("client") or "").strip() or "votre collectivité"
    modele = item.get("modele", "")
    court = re.sub(r"(?i)^smart selection\s+", "", modele).strip()
    mm = re.search(r"(?i)\bR(\d{3})\b", modele)
    if mm: court = "Dell PowerEdge R" + mm.group(1)
    return {"client":client,"qte":item.get("qte","1"),"modele":modele,"modele_court":court,
            "raw_baie":item.get("raw_baie",""),"effective_baie":item.get("effective_baie",""),
            "utile_cluster":item.get("utile_cluster") or item.get("utile",""),
            "capacite":item.get("capacite",""),"nb_vm":item.get("nb_vm","")}

def section_commerciale(doc, arg, item):
    ctx = ctx_item(item)
    add_para(doc, "Présentation & valeur ajoutée", bold=True)
    if arg.get("blocs"):
        for b in arg["blocs"]:
            t = b.get("t")
            if t == "p": add_subst(doc, b.get("x",""), ctx)
            elif t == "h":
                hp = add_para(doc, b.get("x",""), bold=True, size=12)
                for rr in hp.runs: rr.font.color.rgb = ONEID
            elif t == "puces":
                for it in b.get("x", []):
                    p = doc.add_paragraph(); sstyle(doc, p, "Bullet 1"); _fill_runs(p, it, ctx)
            elif t == "img": insert_images(doc, item, arg)
            elif t == "image":
                if not _img_b64(doc, b.get("data"), b.get("width_cm") or 13):
                    insert_images(doc, item, arg)
            elif t == "sous":
                pp = doc.add_paragraph(); sstyle(doc, pp, BODY); _fill_runs(pp, b.get("x",""), ctx)
                for rr in pp.runs: rr.bold = True
            elif t == "table":
                headers = b.get("headers", ["Élément", "Détail"])
                tb = doc.add_table(rows=1, cols=2); style_table(doc, tb)
                set_cell(tb.rows[0].cells[0], headers[0], bold=True); set_cell(tb.rows[0].cells[1], headers[1], bold=True)
                for row in b.get("rows", []):
                    if isinstance(row, (list, tuple)) and len(row) >= 2:
                        c = tb.add_row().cells
                        c[0].text = ""; _fill_runs(c[0].paragraphs[0], str(row[0]), ctx)
                        c[1].text = ""; _fill_runs(c[1].paragraphs[0], str(row[1]), ctx)
            elif t == "vol":
                keys = re.findall(r"\{(\w+)\}", b.get("x",""))
                if keys and all(ctx.get(k) for k in keys): add_subst(doc, b.get("x",""), ctx)
    else:
        if arg.get("presentation"): add_para(doc, arg["presentation"])
        for pf in arg.get("points_forts", []):
            p = doc.add_paragraph(); sstyle(doc, p, "Bullet 1")
            r = p.add_run(pf.get("titre","")); r.bold = True
            if pf.get("detail"): p.add_run(" — " + pf["detail"])
        if arg.get("contexte_marche"): add_para(doc, "Contexte marché : " + arg["contexte_marche"])
    if arg.get("sources"):
        add_para(doc, "Sources : " + " ; ".join(arg["sources"]), italic=True, size=8)
    if arg.get("avertissement"):
        add_para(doc, arg["avertissement"], italic=True, size=8)

def render_item(doc, item, fields, cat_label):
    titre = item.get("role") or ""
    if re.match(r"(?i)^group\s*\d+$", titre.strip()): titre = ""
    modele = item.get("modele", "")
    head = (titre + " — " + modele).strip(" —") or cat_label
    add_heading(doc, head, 3, qte=item.get("qte", "1"))

    arg = match_arg(ARGS, modele)
    if arg and arg.get("titre") not in SHOWN_ARGS:
        section_commerciale(doc, arg, item); SHOWN_ARGS.add(arg.get("titre"))

    rows = [(lbl, item.get(k)) for lbl, k in fields]
    if any(v not in (None,"","À préciser") for _, v in rows):
        add_heading(doc, "Configuration proposée", 3)
        kv_table(doc, rows, red_values=True)
    if item.get("comparatif"):
        add_heading(doc, "Comparatif constructeur", 3)
        # (rendu simple)
        for row in item["comparatif"]:
            add_para(doc, "• " + str(row.get("carac","")) + " : " + str(row.get("retenu","")))
    if item.get("composants"):
        add_para(doc, "Configuration détaillée (nomenclature)", bold=True)
        nomenclature_table(doc, item["composants"])

def section_presentation_oneid(doc):
    """Chapitre permanent de presentation de la societe ONE ID (boilerplate)."""
    add_heading(doc, "Présentation ONE ID", 1)
    add_para(doc, "ONE ID, ce sont 32 personnes expérimentées qui travaillent à apporter à nos "
             "clients les meilleures réponses techniques à leurs problématiques informatiques.")
    add_para(doc, "Nous sommes structurés en 2 équipes opérationnelles :")
    for b in [
        "Système d'Information : en charge des projets logiciels autour du système d'information "
        "(intégration d'ERP, CRM, Décisionnel…) et du développement d'applications métiers "
        "(Windows, Web et Mobile).",
        "Infrastructure : audite, conseille, déploie et maintient les solutions serveurs, "
        "virtualisation, stockage, réseau et sécurité."]:
        bp = doc.add_paragraph(); sstyle(doc, bp, "Bullet 1"); _fill_runs(bp, b, {})
    add_para(doc, "Nous sommes partenaires certifiés Dell / Proxmox / VMware / Veeam / WithSecure / "
             "Fortinet / Microsoft… Nos équipes suivent une formation continue pour être certifiées "
             "sur les produits que nous proposons.")
    add_para(doc, "Une équipe commerciale constituée d'une directrice commerciale, 2 ingénieurs "
             "commerciaux, 2 assistantes ADV et une assistante communication et marketing est à "
             "votre écoute pour répondre à vos demandes.")
    # Champ image ONE ID (fichier 'oneid_presentation.png' a cote du template), sinon emplacement reserve
    img = os.path.join(os.path.dirname(BASE), "oneid_presentation.png")
    if os.path.exists(img):
        try:
            doc.add_picture(img, width=Cm(14))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        except Exception:
            pass
    else:
        ph = add_para(doc, "[ Image ONE ID à insérer ]", red=True, italic=True)
        ph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_para(doc, "La satisfaction client étant notre priorité, nous vérifions tous les ans que "
             "nous soyons en adéquation avec cet objectif via une enquête de satisfaction.")
    sp = add_para(doc, "Nous sommes au-delà des 90 % de satisfaction client chaque année.", bold=True)
    for r in sp.runs: r.font.color.rgb = ONEID


def inventaire_table(doc, spec):
    lignes = []
    for key, label, _intro, _f in CATS:
        for it in spec.get(key, []) or []:
            lignes.append((label, it.get("modele",""), it.get("qte","1")))
    if not lignes: return
    t = doc.add_table(rows=1, cols=3); style_table(doc, t)
    for i, h in enumerate(("Domaine","Équipement proposé","Qté")):
        set_cell(t.rows[0].cells[i], h, bold=True)
    for dom, eq, q in lignes:
        c = t.add_row().cells
        set_cell(c[0], dom); set_cell(c[1], eq, red=True); set_cell(c[2], q, red=True)
    _set_widths(t, [Cm(4.0), Cm(11.0), Cm(1.4)])


def render_solution_text(doc, txt, ctx):
    for line in str(txt).split("\n"):
        line = line.strip()
        if not line:
            continue
        if line[:1] in ("-", "\u2022", "*"):
            bp = doc.add_paragraph(); sstyle(doc, bp, "Bullet 1"); _fill_runs(bp, line[1:].strip(), ctx)
        else:
            add_subst(doc, line, ctx)


def insert_arch_image(doc, aff):
    placed = False
    for im in (aff.get("images_architecture") or []):
        data = im.get("data") if isinstance(im, dict) else im
        if not data: continue
        if "," in data and "base64" in data[:40]: data = data.split(",", 1)[1]
        try:
            doc.add_picture(io.BytesIO(base64.b64decode(data)), width=Cm(16))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER; placed = True
        except Exception: pass
    if not placed:
        ph = add_para(doc, "[ Schéma d'architecture à insérer ]", red=True, italic=True)
        ph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def section_solution_proposee(doc, aff, spec):
    add_heading(doc, "Solution proposée", 1)
    ctx = {"client": (aff.get("client") or "").strip() or "votre collectivité"}
    if aff.get("solution_proposee"):
        render_solution_text(doc, aff["solution_proposee"], ctx)
    add_heading(doc, "Inventaire de la solution", 2)
    add_para(doc, "Vue d'ensemble des équipements proposés dans le cadre de ce dossier :")
    inventaire_table(doc, spec)
    add_heading(doc, "Architecture proposée", 2)
    insert_arch_image(doc, aff)


def _shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill); tcPr.append(shd)


def _gcell(cell, text, size=8, bold=False, color=None, center=False):
    cell.text = ""; p = cell.paragraphs[0]
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("" if text is None else str(text)); r.font.size = Pt(size); r.bold = bold
    if color:
        r.font.color.rgb = RGBColor.from_string(color)


def _pdate(s):
    try:
        return datetime.date.fromisoformat((s or "").strip())
    except Exception:
        return None


MONTHS_FR = ["janv.", "févr.", "mars", "avr.", "mai", "juin",
             "juil.", "août", "sept.", "oct.", "nov.", "déc."]


def _fmtd(s):
    d = _pdate(s)
    return d.strftime("%d/%m/%Y") if d else (s or "")


def _set_widths(table, widths):
    table.autofit = False
    table.allow_autofit = False
    tblPr = table._tbl.tblPr
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            if i < len(widths):
                cell.width = widths[i]


def _gantt_buf(tasks):
    """Gantt hebdomadaire en grille (image PNG) calqué sur l'aperçu HTML ; None si indisponible."""
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        from matplotlib.patches import Rectangle
    except Exception:
        return None
    n = len(tasks)
    if n == 0:
        return None
    ds = []
    for tk in tasks:
        d0, d1 = _pdate(tk.get("debut")), _pdate(tk.get("fin"))
        if d0: ds.append(d0)
        if d1: ds.append(d1)
    if not ds:
        return None
    monday = lambda d: d - datetime.timedelta(days=d.weekday())
    mn, mx = monday(min(ds)), max(ds)
    weeks, cur = [], mn
    while cur <= mx and len(weeks) < 80:
        weeks.append(cur); cur = cur + datetime.timedelta(days=7)
    nW = len(weeks)

    # Couleurs (alignées sur l'aperçu HTML)
    NAVY = ("#" + HDR_HEX) if DESIGN_ON else "#21425f"
    ORANGE, PURPLE, GRID = "#f4b183", "#c9b7e4", "#dfe3e8"
    RECEP, TELE = "#8bc98b", "#6aaed6"   # réception matériel (vert), télétravail (bleu)
    # Largeurs de colonnes en unités arbitraires
    ID_W, TASK_W, WK_W, ROW_H = 0.7, 7.0, 0.95, 1.0
    total_w = ID_W + TASK_W + WK_W * nW
    total_h = ROW_H * (n + 1)  # +1 ligne d'en-tête
    fig_w = max(7.0, min(11.0, total_w * 0.17))
    fig_h = max(1.8, min(15.0, total_h * 0.17))
    fig, ax = plt.subplots(figsize=(fig_w, fig_h))
    ax.set_xlim(0, total_w); ax.set_ylim(0, total_h); ax.invert_yaxis(); ax.axis("off")
    xcol = lambda i: ID_W + TASK_W + WK_W * i

    # Quadrillage léger (dessous)
    for i in range(nW + 1):
        ax.plot([xcol(i)] * 2, [0, total_h], color=GRID, lw=0.4, zorder=0)
    ax.plot([ID_W] * 2, [0, total_h], color=GRID, lw=0.4, zorder=0)
    for r in range(n + 2):
        ax.plot([0, total_w], [ROW_H * r] * 2, color=GRID, lw=0.4, zorder=0)

    # En-tête bleu marine
    ax.add_patch(Rectangle((0, 0), total_w, ROW_H, facecolor=NAVY, edgecolor=NAVY, zorder=1))
    ax.text(ID_W / 2, ROW_H / 2, "ID", ha="center", va="center", color="white", fontsize=7, fontweight="bold", zorder=2)
    ax.text(ID_W + 0.12, ROW_H / 2, "Tâche", ha="left", va="center", color="white", fontsize=7, fontweight="bold", zorder=2)
    for i, w in enumerate(weeks):
        ax.text(xcol(i) + WK_W / 2, ROW_H / 2, w.strftime("%d/%m"),
                ha="center", va="center", color="white", fontsize=5.5, zorder=2)

    # Lignes de tâches
    for r, tk in enumerate(tasks):
        y = ROW_H * (r + 1)
        ax.text(ID_W / 2, y + ROW_H / 2, str(tk.get("id") or (r + 1)),
                ha="center", va="center", fontsize=6, color="#222", zorder=2)
        nm = tk.get("nom") or ""
        if len(nm) > 60: nm = nm[:58] + "…"
        ax.text(ID_W + 0.12, y + ROW_H / 2, nm, ha="left", va="center", fontsize=6, color="#222", zorder=2)
        d0, d1 = _pdate(tk.get("debut")), _pdate(tk.get("fin"))
        if d0 and d1:
            if tk.get("conge"):
                color = PURPLE
            elif re.search(r"r[ée]ception", tk.get("nom") or "", re.I):
                color = RECEP
            elif tk.get("teletravail"):
                color = TELE
            else:
                color = ORANGE
            for i, w in enumerate(weeks):
                we = w + datetime.timedelta(days=6)
                if d0 <= we and d1 >= w:
                    ax.add_patch(Rectangle((xcol(i), y), WK_W, ROW_H,
                                           facecolor=color, edgecolor="white", lw=0.6, zorder=1))

    fig.tight_layout(pad=0.2)
    buf = io.BytesIO(); fig.savefig(buf, format="png", dpi=200, bbox_inches="tight")
    plt.close(fig); buf.seek(0)
    return buf


def section_planning(doc, spec):
    tasks = [t for t in (spec.get("planning") or []) if (t.get("nom") or "").strip()]
    if not tasks:
        return
    ONEID_HX = HDR_HEX
    add_heading(doc, "Planning prévisionnel infra", 1)
    add_para(doc, "Planning prévisionnel des interventions. Le diagramme ci-dessous présente "
                  "l'enchaînement hebdomadaire des tâches — orange : tâche ; vert : réception matériel ; "
                  "bleu : télétravail ; violet : congé.", size=9)
    info = ["ID", "Nom de tâche", "Début", "Fin", "Durée", "Équipes", "Tél."]
    t = doc.add_table(rows=1, cols=len(info))
    try:
        t.style = doc.styles["Table Grid"]
    except KeyError:
        pass
    widths = [Cm(x) for x in (0.9, 5.5, 2.2, 2.2, 1.2, 2.6, 0.8)]
    hdr = t.rows[0].cells
    for i, h in enumerate(info):
        _gcell(hdr[i], h, 8, True, "ffffff"); _shade(hdr[i], ONEID_HX)
    for k, tk in enumerate(tasks, 1):
        c = t.add_row().cells
        _gcell(c[0], tk.get("id") or k, 8, center=True)
        _gcell(c[1], tk.get("nom", ""), 8)
        _gcell(c[2], _fmtd(tk.get("debut")), 8, center=True)
        _gcell(c[3], _fmtd(tk.get("fin")), 8, center=True)
        _gcell(c[4], tk.get("duree", ""), 8, center=True)
        _gcell(c[5], tk.get("equipes", ""), 8)
        _gcell(c[6], tk.get("teletravail", ""), 8, center=True)
    _set_widths(t, widths)

    buf = _gantt_buf(tasks)
    if buf is not None:
        add_para(doc, "")
        doc.add_picture(buf, width=Cm(16.5))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


FR_NUM = {1: "une", 2: "deux", 3: "trois", 4: "quatre", 5: "cinq", 6: "six", 7: "sept", 8: "huit"}


def _bullet(doc, text, ctx=None):
    bp = doc.add_paragraph(); sstyle(doc, bp, "Bullet 1"); _fill_runs(bp, text, ctx or {})
    return bp


def _img_ph(doc, label="[ Schéma à insérer ]"):
    p = add_para(doc, label, red=True, italic=True)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def _img_b64(doc, data, width_cm=13):
    """Insère une image encodée en base64 (data-URL ou brute), centrée. True si insérée."""
    if not data:
        return False
    if "," in data and "base64" in data[:40]:
        data = data.split(",", 1)[1]
    try:
        doc.add_picture(io.BytesIO(base64.b64decode(data)), width=Cm(width_cm))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        return True
    except Exception:
        return False


def _load_arg(folder):
    """Charge un argumentaire éditable Documentation_Constructeur/<folder>/argumentaire.json."""
    base = DOC_BASE or os.path.join(os.getcwd(), "Documentation_Constructeur")
    fp = os.path.join(base, folder, "argumentaire.json")
    if os.path.exists(fp):
        try:
            return json.load(open(fp, encoding="utf-8"))
        except Exception:
            return None
    return None


def render_blocs(doc, blocs, ctx, conditions=None):
    """Rendu commun des blocs d'argumentaire : p / h / sous / puces / img / table.
    Un bloc peut porter une clé 'if' (nom de condition) ; il n'est rendu que si la condition est vraie."""
    conditions = conditions or {}
    for b in blocs or []:
        cond = b.get("if")
        if cond and not conditions.get(cond):
            continue
        t = b.get("t")
        if t == "p":
            add_subst(doc, b.get("x", ""), ctx)
        elif t in ("h", "h2"):
            add_heading(doc, b.get("x", ""), 3)
        elif t == "sous":
            p = doc.add_paragraph(); sstyle(doc, p, BODY)
            _fill_runs(p, b.get("x", ""), ctx)
            for r in p.runs:
                r.bold = True
        elif t == "puces":
            for it in b.get("x", []):
                _bullet(doc, it, ctx)
        elif t == "img":
            _img_ph(doc, b.get("x", "[ Schéma à insérer ]"))
        elif t == "image":
            if not _img_b64(doc, b.get("data"), b.get("width_cm") or 13):
                _img_ph(doc, "[ Image à insérer ]")
        elif t == "table":
            headers = b.get("headers", ["Élément", "Détail"])
            tb = doc.add_table(rows=1, cols=2); style_table(doc, tb)
            set_cell(tb.rows[0].cells[0], headers[0], bold=True)
            set_cell(tb.rows[0].cells[1], headers[1], bold=True)
            for row in b.get("rows", []):
                if not isinstance(row, (list, tuple)) or len(row) < 2:
                    continue
                c = tb.add_row().cells
                c[0].text = ""; _fill_runs(c[0].paragraphs[0], str(row[0]), ctx)
                c[1].text = ""; _fill_runs(c[1].paragraphs[0], str(row[1]), ctx)


def section_sauvegarde_combo(doc, spec):
    """Chapitre auto piloté par Documentation_Constructeur/SAUVEGARDE_PS_PP/argumentaire.json.
    Inséré uniquement si tous les mots-clés de 'trigger.all_of' sont présents dans la config."""
    arg = _load_arg("SAUVEGARDE_PS_PP")
    if not arg:
        return
    allitems = []
    for key in ("stockages", "sauvegardes", "serveurs", "switches", "logiciels_services"):
        allitems += (spec.get(key) or [])

    def detect(kw):
        out = []
        for it in allitems:
            if not isinstance(it, dict):
                continue
            s = " ".join(str(it.get(f, "")) for f in ("modele", "role", "designation", "nom", "reference")).lower()
            if kw in s:
                out.append(it)
        return out

    keywords = (arg.get("trigger") or {}).get("all_of") or ["powerstore", "powerprotect"]
    matched = {kw: detect(kw.lower()) for kw in keywords}
    if not all(matched.values()):
        return

    def qty(items):
        n = 0
        for it in items:
            try:
                n += int(str(it.get("qte") or 1))
            except Exception:
                n += 1
        return n or len(items)

    def model(items, default):
        for it in items:
            m = (it.get("modele") or it.get("designation") or it.get("role") or "").strip()
            if m:
                return m
        return default

    ps = matched.get("powerstore") or []
    pp = matched.get("powerprotect") or []
    nps = qty(ps) if ps else 1
    aff = spec.get("affaire", {})
    ctx = {
        "client": (aff.get("client") or "").strip() or "votre collectivité",
        "nps": FR_NUM.get(nps, str(nps)),
        "ps": model(ps, "PowerStore"),
        "pp": model(pp, "PowerProtect"),
    }
    # Conditions d'affichage des blocs "if" : drapeau utilisateur (onglet) OU détection dans la conf
    conditions = {
        "veeam": bool(aff.get("sauv_veeam")) or bool(detect("veeam")),
        "datadomain": bool(aff.get("sauv_datadomain")) or bool(detect("datadomain")) or bool(detect("data domain")),
    }

    add_heading(doc, arg.get("titre", "Solution de Sauvegarde"), 2)
    render_blocs(doc, arg.get("blocs", []), ctx, conditions)
    if arg.get("sources"):
        add_para(doc, "Sources : " + " ; ".join(arg["sources"]), italic=True, size=8)
    if arg.get("avertissement"):
        add_para(doc, arg["avertissement"], italic=True, size=8)


SAUV_COLS = ["VM Fichiers", "VM Database", "VM", "NAS Fichiers", "Physique Database"]
SAUV_ROWS = ["TB", "Daily Change", "Rétentions", "Croissance estimée", "NB de Serveurs"]
SAUV_SUM_ROWS = {"TB", "NB de Serveurs"}


def _volumetrie_table(doc, site):
    nom = (site.get("nom") or "Site").strip()
    cells = site.get("cells") or {}
    add_para(doc, nom, bold=True)
    t = doc.add_table(rows=1, cols=len(SAUV_COLS) + 2); style_table(doc, t)
    hdr = t.rows[0].cells
    _gcell(hdr[0], "Typologie des Données", 9, True, "ffffff"); _shade(hdr[0], HDR_HEX)
    for i, c in enumerate(SAUV_COLS):
        _gcell(hdr[i + 1], c, 9, True, "ffffff"); _shade(hdr[i + 1], HDR_HEX)
    _gcell(hdr[-1], "Total", 9, True, "ffffff"); _shade(hdr[-1], HDR_HEX)
    for row in SAUV_ROWS:
        rc = t.add_row().cells
        _gcell(rc[0], row, 9, True)
        rowvals = cells.get(row, {}) if isinstance(cells.get(row), dict) else {}
        total = 0.0; has = False
        for i, c in enumerate(SAUV_COLS):
            v = str(rowvals.get(c, "") or "")
            _gcell(rc[i + 1], v, 9, center=True)
            if row in SAUV_SUM_ROWS:
                try:
                    total += float(v.replace(",", ".").strip()); has = True
                except Exception:
                    pass
        tot = ""
        if row in SAUV_SUM_ROWS and has:
            tot = str(int(total)) if total == int(total) else str(total)
        _gcell(rc[-1], tot, 9, True, center=True)
    _set_widths(t, [Cm(3.6)] + [Cm(2.3)] * len(SAUV_COLS) + [Cm(1.4)])


def section_sauvegarde_offre(doc, spec):
    """Chapitre Sauvegarde piloté par l'onglet Fonctionnalité : volumétrie + argumentaire du cas."""
    sv = spec.get("sauvegarde") or {}
    log = (sv.get("logiciel") or "").lower()
    cib = (sv.get("cible") or "").lower()
    sites = [s for s in (sv.get("sites") or []) if isinstance(s, dict)]
    if not (log and cib) and not sites:
        return
    cl = (spec.get("affaire", {}).get("client") or "").strip() or "votre collectivité"
    add_heading(doc, "Solution de sauvegarde", 2)
    if sites:
        add_para(doc, "Volumétrie et typologie des données à protéger :", bold=True)
        for s in sites:
            _volumetrie_table(doc, s)
            add_para(doc, "")
    folder = {"veeam": "SAUV_VEEAM", "ppdm": "SAUV_PPDM"}.get(log)
    code = {"san": "SAN", "dd": "DD"}.get(cib)
    if folder and code:
        case = _load_arg(folder + "_" + code)
        if case:
            lic = (sv.get("licence") or "").lower()
            conditions = {"licence_acheter": lic == "acheter", "licence_deja": lic == "deja",
                          "cloud": bool(sv.get("cloud"))}
            render_blocs(doc, case.get("blocs", []), {"client": cl}, conditions)
            if case.get("sources"):
                add_para(doc, "Sources : " + " ; ".join(case["sources"]), italic=True, size=8)
            if case.get("avertissement"):
                add_para(doc, case["avertissement"], italic=True, size=8)


def _append_admin(doc):
    """Ajoute les chapitres administratifs figés (assets/CHAPITRES_ADMIN.docx) à la fin du document.
    Retourne le Composer (pour la sauvegarde) ou None si indisponible."""
    path = os.path.join(os.path.dirname(BASE), "CHAPITRES_ADMIN.docx")
    if not os.path.exists(path):
        alt = os.path.join(HERE, "assets", "CHAPITRES_ADMIN.docx")
        if os.path.exists(alt):
            path = alt
    if not os.path.exists(path):
        return None
    try:
        from docxcompose.composer import Composer
        doc.add_page_break()
        comp = Composer(doc)
        comp.append(Document(path))
        return comp
    except Exception:
        return None


def section_prestations(doc, spec):
    """Chapitre Prestations et méthodologie : intitulé + durée (rouge) + méthodologie."""
    prs = [p for p in (spec.get("prestations") or [])
           if isinstance(p, dict) and (p.get("intitule") or "").strip()]
    if not prs:
        return
    cl = (spec.get("affaire", {}).get("client") or "").strip() or "votre collectivité"
    add_heading(doc, "Prestations et méthodologie", 1)
    for p in prs:
        # 1) Intitulé
        add_heading(doc, (p.get("intitule") or "").strip(), 2)
        # 2) Argumentaire lié (sélectionné dans l'onglet Prestations)
        folder = (p.get("argumentaire") or "").strip()
        if folder:
            arg = _load_arg(folder)
            if arg:
                render_blocs(doc, arg.get("blocs", []), {"client": cl}, {})
                if arg.get("sources"):
                    add_para(doc, "Sources : " + " ; ".join(arg["sources"]), italic=True, size=8)
        # 3) Complément / options
        comp = p.get("complement") or p.get("methodo") or ""
        if comp.strip():
            render_solution_text(doc, comp, {"client": cl})
        # 4) Durée estimée (en dernier)
        duree = (p.get("duree") or "").strip()
        if duree:
            add_para(doc, "Durée estimée : " + duree, bold=True, red=True)


def section_chapitres(doc, spec):
    """Chapitres rédigés issus du Kanban : titre + texte éditable + points à traiter (puces rouges).
    Rendus dans l'ordre défini dans le Kanban."""
    chaps = spec.get("chapitres") or []
    cl = (spec.get("affaire", {}).get("client") or "").strip() or "votre collectivité"
    for ch in chaps:
        if not isinstance(ch, dict):
            continue
        titre = (ch.get("titre") or "").strip()
        if not titre:
            continue
        add_heading(doc, titre, 1)
        # 1) Argumentaire lié (sélectionné dans la description du chapitre)
        folder = (ch.get("argumentaire") or "").strip()
        if folder:
            arg = _load_arg(folder)
            if arg:
                render_blocs(doc, arg.get("blocs", []), {"client": cl}, {})
                if arg.get("sources"):
                    add_para(doc, "Sources : " + " ; ".join(arg["sources"]), italic=True, size=8)
        # 2) Complément d'info / options / prérequis (ou ancien champ « texte »)
        texte = ch.get("complement") or ch.get("texte") or ""
        if texte.strip():
            render_solution_text(doc, texte, {"client": cl})
        for im in (ch.get("images") or []):
            data = im.get("data") if isinstance(im, dict) else im
            if not data:
                continue
            if "," in data and "base64" in data[:40]:
                data = data.split(",", 1)[1]
            try:
                doc.add_picture(io.BytesIO(base64.b64decode(data)), width=Cm(15))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            except Exception:
                pass
        cartes = []
        for c in (ch.get("cartes") or []):
            if isinstance(c, dict):
                lab = (c.get("l") or c.get("label") or "").strip(); ok = bool(c.get("ok"))
            else:
                lab = str(c).strip(); ok = False
            if lab:
                cartes.append((lab, ok))
        if cartes:
            add_para(doc, "Points à traiter :", bold=True)
            for lab, ok in cartes:
                p = doc.add_paragraph(); sstyle(doc, p, "Bullet 1")
                r = p.add_run(("✔ " if ok else "") + lab)
                r.font.color.rgb = RGBColor(0x1E, 0x7E, 0x34) if ok else RED   # vert = validé, rouge = à traiter
                _afont(r)
        # 3) Durée d'installation estimée (en dernier)
        duree = (ch.get("duree") or "").strip()
        if duree:
            add_para(doc, "Durée d'installation estimée : " + duree, bold=True, red=True)


def section_pra_pca(doc, aff):
    """Chapitre PRA/PCA si activé : tableau paramètres (valeurs en rouge) + argumentaire."""
    pp = aff.get("pra_pca") or {}
    if not pp.get("actif"):
        return
    cl = (aff.get("client") or "").strip() or "votre collectivité"
    mode = (pp.get("mode") or "").lower()
    titre_mode = "Plan de Continuité d'Activité (PCA)" if mode == "pca" else "Plan de Reprise d'Activité (PRA)"
    CIBLES = {"ancien": "Ancien matériel reconditionné (site de repli)",
              "nouveau": "Nouveau matériel dédié", "cloud": "Cloud ONE ID"}
    METHODES = {"veeam": "Veeam (logiciel, indépendant du matériel)",
                "natif": "Réplication native de la baie de stockage"}
    add_heading(doc, titre_mode, 2)
    info = [("Dispositif", titre_mode),
            ("RTO (délai de reprise visé)", pp.get("rto", "")),
            ("RPO (perte de données max admise)", pp.get("rpo", "")),
            ("Cible de réplication", CIBLES.get(pp.get("cible"), pp.get("cible", ""))),
            ("Méthode de réplication", METHODES.get(pp.get("methode"), pp.get("methode", "")))]
    kv_table(doc, info, headers=("Paramètre", "Valeur"), red_values=True, widths=[Cm(6.5), Cm(9.5)])
    add_para(doc, "")
    texte = pp.get("texte") or ""
    if texte.strip():
        render_solution_text(doc, texte, {"client": cl})
    # Cas détaillés éditables (Documentation_Constructeur/<dossier>/argumentaire.json),
    # sélectionnés selon (dispositif, cible, méthode).
    CASE_FILES = {
        ("pra", "cloud", "veeam"): "PRA_VEEAM",            # offre détaillée fournie
        ("pra", "ancien", "veeam"): "PRA_ANCIEN_VEEAM",
        ("pra", "cloud", "natif"): "PRA_CLOUD_NATIF",
        ("pra", "nouveau", "veeam"): "PRA_NOUVEAU_VEEAM",
        ("pra", "nouveau", "natif"): "PRA_NOUVEAU_NATIF",
        ("pra", "ancien", "natif"): "PRA_ANCIEN_NATIF",
        ("pca", "cloud", "veeam"): "PCA_CLOUD_VEEAM",
        ("pca", "cloud", "natif"): "PCA_CLOUD_NATIF",
        ("pca", "nouveau", "veeam"): "PCA_NOUVEAU_VEEAM",
        ("pca", "nouveau", "natif"): "PCA_NOUVEAU_NATIF",
        ("pca", "ancien", "veeam"): "PCA_ANCIEN_VEEAM",
        ("pca", "ancien", "natif"): "PCA_ANCIEN_NATIF",
    }
    key = (mode, (pp.get("cible") or "").lower(), (pp.get("methode") or "").lower())
    folder = CASE_FILES.get(key)
    if folder:
        case = _load_arg(folder)
        if case:
            ctxc = {"client": cl, "rto": pp.get("rto") or "à préciser", "rpo": pp.get("rpo") or "à préciser"}
            render_blocs(doc, case.get("blocs", []), ctxc)
            if case.get("avertissement"):
                add_para(doc, case["avertissement"], italic=True, size=8)
    add_para(doc, "Note : les valeurs RTO/RPO et les modalités de réplication sont propres à ce dossier "
                  "et à valider techniquement selon la configuration retenue.", italic=True, size=9)


def section_verification(doc, spec):
    """Vérification de conformité CCTP vs solution, en couleur de relecture (à supprimer avant envoi)."""
    items = [v for v in (spec.get("verification") or [])
             if isinstance(v, dict) and (v.get("exigence") or v.get("suggestion") or v.get("theme"))]
    if not items:
        return
    REVIEW = "7030A0"   # violet de relecture, distinct du rouge (champs personnalisables)
    LBL = {"couvert": "Couvert", "ecart": "Écart", "a_preciser": "À préciser"}
    add_heading(doc, "Vérification de conformité au CCTP (à valider)", 1)
    add_para(doc, "Note de relecture : section générée automatiquement pour contrôler l'adéquation "
                  "de la solution au CCTP. À vérifier puis SUPPRIMER avant envoi au client.",
             color=REVIEW, italic=True, size=9)
    for v in items:
        st = v.get("statut") or "a_preciser"
        theme = (v.get("theme") or "").strip()
        head = LBL.get(st, "À préciser") + ((" — " + theme) if theme else "")
        body = (v.get("suggestion") or v.get("exigence") or "").strip()
        p = doc.add_paragraph(); sstyle(doc, p, BODY)
        r = p.add_run("• " + head + (" : " if body else "")); r.bold = True
        r.font.color.rgb = RGBColor.from_string(REVIEW)
        if body:
            r2 = p.add_run(body); r2.font.color.rgb = RGBColor.from_string(REVIEW)
        if v.get("exigence") and v.get("suggestion"):
            add_para(doc, "Exigence CCTP : " + str(v["exigence"]), color=REVIEW, italic=True, size=9)


def patch_cover(path, aff):
    d = aff.get("date","")
    try:
        dt = datetime.date.fromisoformat(d); ma = MOIS[dt.month-1] + " " + str(dt.year)
    except Exception: ma = ""
    repl = {"Titre principal": (aff.get("projet") or aff.get("client") or "Document technique"),
            "Sous-titre document": "Dossier technique — Solution d'infrastructure",
            "Type de doc": "Réponse à appel d'offre", "Mois année": ma}
    tmp = path + ".tmp"
    with zipfile.ZipFile(path,"r") as zin, zipfile.ZipFile(tmp,"w",zipfile.ZIP_DEFLATED) as zout:
        for it in zin.infolist():
            data = zin.read(it.filename)
            if it.filename in ("word/document.xml","word/header2.xml"):
                txt = data.decode("utf-8")
                for k, v in repl.items():
                    if v: txt = txt.replace(k, v)
                if aff.get("version"): txt = txt.replace("Version ", "Version " + str(aff["version"]) + " ")
                data = txt.encode("utf-8")
            zout.writestr(it, data)
    shutil.move(tmp, path)

def main():
    global ARGS, AFF, DOC_BASE
    spec = json.load(open(sys.argv[1], encoding="utf-8")); out_path = sys.argv[2]
    aff = spec.get("affaire", {}); AFF = aff
    if len(sys.argv) > 3: DOC_BASE = sys.argv[3]
    elif not DOC_BASE: DOC_BASE = os.environ.get("DOC_BASE") or os.path.join(HERE, "..", "Documentation_Constructeur")
    if not os.path.isdir(DOC_BASE): DOC_BASE = os.path.join(os.getcwd(), "Documentation_Constructeur")
    ARGS = load_argumentaires(); SHOWN_ARGS.clear()
    cl = (aff.get("client") or "").strip() or "votre collectivité"
    actx = {"client":cl,"projet":aff.get("projet",""),"reference":aff.get("reference",""),
            "prix":aff.get("prix_solution","")}

    apply_theme(spec.get("_theme"))
    doc = Document(BASE)
    # Intro engageante (réécrite par le mode Design si présent)
    DEFAULT_INTRO = ("ONE ID a le plaisir de vous présenter la solution d'infrastructure conçue "
                     "pour {client} dans le cadre du projet {projet}. Ce document détaille les équipements "
                     "proposés, leurs caractéristiques techniques et la valeur apportée à votre organisation.")
    render_solution_text(doc, aff.get("intro_override") or DEFAULT_INTRO, actx)
    # Bloc affaire (table propre, valeurs en rouge)
    info = [("Client", actx["client"]), ("Projet", actx["projet"]),
            ("Référence", actx["reference"]), ("Prix catalogue de la solution", actx["prix"]),
            ("Date", aff.get("date","")), ("Version", aff.get("version",""))]
    kv_table(doc, info, headers=("Informations du dossier","Valeur"), red_values=True, widths=[Cm(5.0), Cm(11.4)])
    # Legende relecture
    add_para(doc, "Note de relecture : les éléments affichés en rouge sont personnalisés "
             "(spécifiques à ce dossier) et sont à vérifier avant envoi au client.", red=True, italic=True, size=9)
    section_presentation_oneid(doc)
    contexte = aff.get("contexte") or aff.get("notes")
    if contexte:
        add_heading(doc, "Contexte", 2)
        render_solution_text(doc, contexte, {"client": cl})
    section_solution_proposee(doc, aff, spec)

    for key, label, intro, fields in CATS:
        items = spec.get(key) or []
        if not items: continue
        add_heading(doc, label, 2)
        if intro: add_para(doc, intro)
        for it in items:
            render_item(doc, it, fields, label)

    section_pra_pca(doc, aff)
    section_sauvegarde_combo(doc, spec)
    section_sauvegarde_offre(doc, spec)
    section_chapitres(doc, spec)
    section_prestations(doc, spec)
    section_planning(doc, spec)
    composer = _append_admin(doc)        # chapitres administratifs figés, à la fin
    section_verification(doc, spec)
    (composer or doc).save(out_path)
    patch_cover(out_path, aff); print("OK ->", out_path)

if __name__ == "__main__":
    main()

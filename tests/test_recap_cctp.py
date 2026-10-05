# /home/numa/projets/appel-offre/tests/test_recap_cctp.py
"""Tests de l'étape « Récapitulatif du CCTP — Focus technique ».

Couvre : parsing/normalisation de la sortie LLM (recap_cctp), cohérence des colonnes
entre le backend et index.html, rendu Word (engine/generer_doc.section_recap_cctp)
et endpoint /api/recap-cctp avec LLM simulé (ignoré si fastapi/httpx absents).

Exécution (depuis la racine du dépôt) : python -m unittest tests.test_recap_cctp -v
"""
from __future__ import annotations

import json
import os
import re
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "app"))
sys.path.insert(0, str(ROOT / "app" / "engine"))

import recap_cctp as rc  # noqa: E402

LIGNE_MATRICE = {"domaine": "Sauvegarde", "exigences": "Sauvegarde quotidienne.", "preuves": "Fiche",
                 "capacite_oneid": "À confirmer : référence", "risque": "Moyen : volumétrie",
                 "action": "Demander la volumétrie", "ref": "CCTP §4.2", "statut": "a_qualifier"}


class ParseRecapTests(unittest.TestCase):
    def test_json_entoure_de_fence_et_think(self) -> None:
        raw = "<think>blabla</think>\n```json\n" + json.dumps({"matrice": [LIGNE_MATRICE]}) + "\n```"
        r = rc.parse_recap(raw)
        self.assertTrue(r["ok"])
        self.assertEqual(r["matrice"][0]["ref"], "CCTP §4.2")
        self.assertEqual(r["plan"], [])
        self.assertEqual(r["ressources"], [])

    def test_texte_autour_du_json(self) -> None:
        raw = "Voici le résultat : " + json.dumps({"matrice": [LIGNE_MATRICE]}) + " Fin."
        self.assertTrue(rc.parse_recap(raw)["ok"])

    def test_sortie_inexploitable(self) -> None:
        r = rc.parse_recap("désolé, je ne peux pas")
        self.assertFalse(r["ok"])
        self.assertEqual([r[k] for k in rc.TABLES], [[], [], []])

    def test_valide_interdit_au_llm(self) -> None:
        row = dict(LIGNE_MATRICE, statut="Validé")
        r = rc.parse_recap(json.dumps({"matrice": [row]}))
        self.assertEqual(r["matrice"][0]["statut"], rc.STATUT_DEFAUT)

    def test_valide_autorise_pour_humain(self) -> None:
        row = dict(LIGNE_MATRICE, statut="valide")
        r = rc.normalize_recap({"matrice": [row]}, allow_valide=True)
        self.assertEqual(r["matrice"][0]["statut"], "valide")

    def test_libelles_de_statut_reconnus(self) -> None:
        self.assertEqual(rc.normalize_statut("À préparer"), "a_preparer")
        self.assertEqual(rc.normalize_statut("A Qualifier"), "a_qualifier")
        self.assertEqual(rc.normalize_statut("n'importe quoi"), rc.STATUT_DEFAUT)

    def test_cellules_vides_completees_et_aplaties(self) -> None:
        row = {"role": "Chef de projet", "responsabilite": "", "competences": ["ITIL", "PMP"],
               "presence": "sur\nsite", "go_nogo": "nogo", "statut": "a_confirmer"}
        r = rc.normalize_recap({"ressources": [row]})["ressources"][0]
        self.assertEqual(r["responsabilite"], rc.A_PRECISER)
        self.assertEqual(r["competences"], "ITIL ; PMP")
        self.assertEqual(r["presence"], "sur site")
        self.assertEqual(r["go_nogo"], "No Go")
        self.assertEqual(rc.normalize_recap({"ressources": [dict(row, go_nogo="peut-être")]})
                         ["ressources"][0]["go_nogo"], "À arbitrer")
        self.assertEqual(set(r), {c for c, _ in rc.TABLES["ressources"]["colonnes"]})

    def test_lignes_vides_et_non_dict_ecartees_et_plafond(self) -> None:
        rows = [{"statut": "a_preparer"}, "texte", None] + [dict(LIGNE_MATRICE)] * (rc.MAX_ROWS + 5)
        r = rc.normalize_recap({"matrice": rows})
        self.assertEqual(len(r["matrice"]), rc.MAX_ROWS)

    def test_cellule_tronquee(self) -> None:
        r = rc.normalize_recap({"matrice": [dict(LIGNE_MATRICE, action="x" * 2000)]})
        self.assertLessEqual(len(r["matrice"][0]["action"]), rc.MAX_CELL)

    def test_message_utilisateur(self) -> None:
        m = rc.build_user_message("C" * 70000, [{"label": "RPO 15 min"}, "Licences"], ["Volumétrie ?"], "2x R660")
        self.assertIn("POINTS D'ATTENTION DÉJÀ IDENTIFIÉS :\n- RPO 15 min\n- Licences", m)
        self.assertIn("CLARIFICATIONS :\n- Volumétrie ?", m)
        self.assertIn("SOLUTION ACTUELLEMENT CONFIGURÉE :\n2x R660", m)
        self.assertIn("C" * 60000, m)
        self.assertNotIn("C" * 60001, m)
        self.assertNotIn("POINTS", rc.build_user_message("cctp"))

    def test_prompt_conforme_au_contrat(self) -> None:
        for spec in rc.TABLES.values():
            for col, _ in spec["colonnes"]:
                self.assertIn('"%s"' % col, rc.RECAP_PROMPT)
        self.assertIn("a_qualifier", rc.RECAP_PROMPT)


class FrontendCoherenceTests(unittest.TestCase):
    def test_colonnes_identiques_backend_et_ui(self) -> None:
        html = (ROOT / "app" / "web" / "index.html").read_text(encoding="utf-8")
        self.assertIn("const STEP_ORDER=['cctp','recap',", html)
        self.assertIn("recap_cctp:recap", html)
        block = html[html.index("const RECAP_TABLES="):html.index("const RECAP_STATUTS=")]
        for key, spec in rc.TABLES.items():
            m = re.search(r"\{key:'%s',titre:'([^']*)',cols:\[(.*?)\]\}" % key, block)
            self.assertIsNotNone(m, key)
            self.assertEqual(m.group(1), spec["titre"])
            ui_cols = re.findall(r"\['(\w+)',", m.group(2))
            self.assertEqual(ui_cols, [c for c, _ in spec["colonnes"]])
        self.assertIn("Capacité ONEID à confirmer", block)
        for k, lbl in rc.STATUTS.items():
            self.assertIn("['%s','%s']" % (k, lbl), html)

    def test_colonnes_elargies_identiques_ui_et_word(self) -> None:
        html = (ROOT / "app" / "web" / "index.html").read_text(encoding="utf-8")
        m = re.search(r"const RECAP_WIDE=\[(.*?)\];", html)
        self.assertIsNotNone(m)
        self.assertEqual(tuple(re.findall(r"'(\w+)'", m.group(1))), rc.WIDE_COLS)
        toutes = {c for spec in rc.TABLES.values() for c, _ in spec["colonnes"]}
        self.assertTrue(set(rc.WIDE_COLS) <= toutes)

    def test_affichage_pleine_largeur(self) -> None:
        html = (ROOT / "app" / "web" / "index.html").read_text(encoding="utf-8")
        self.assertNotIn(".wrap{max-width:1240px", html)
        self.assertIn(".wrap{max-width:none;width:100%", html)


class WordExportTests(unittest.TestCase):
    def test_section_recap_dans_le_docx(self) -> None:
        try:
            from docx import Document
            import generer_doc as gd
        except ImportError as e:  # python-docx absent hors conteneur
            self.skipTest("python-docx indisponible : %s" % e)
        doc = Document()
        spec = {"affaire": {"recap_cctp": {
            "matrice": [dict(LIGNE_MATRICE, statut="valide")],
            "plan": [{"periode": "Phase 0", "objectifs": "Lever les ambiguïtés", "statut": "a_preparer"}],
            "ressources": [{"role": "Chef de projet", "go_nogo": "Go", "statut": "a_confirmer"}]}}}
        n_sections = len(doc.sections)
        gd.section_recap_cctp(doc, spec)
        gd._landscape_close(doc)
        self.assertEqual(len(doc.tables), 3)
        self.assertEqual([c.text for c in doc.tables[0].rows[0].cells],
                         [lbl for _, lbl in rc.TABLES["matrice"]["colonnes"]])
        self.assertEqual(doc.tables[0].rows[1].cells[-1].text, "Validé")
        self.assertEqual(doc.tables[2].rows[1].cells[6].text, "Go")
        self.assertEqual(len(doc.sections), n_sections + 2)
        land, port = doc.sections[-2], doc.sections[-1]
        self.assertGreater(land.page_width, land.page_height)
        self.assertLess(port.page_width, port.page_height)
        texts = [p.text for p in doc.paragraphs]
        self.assertIn("Récapitulatif du CCTP — Focus technique", texts)

    def test_section_absente_si_vide(self) -> None:
        try:
            from docx import Document
            import generer_doc as gd
        except ImportError as e:
            self.skipTest("python-docx indisponible : %s" % e)
        doc = Document()
        gd.section_recap_cctp(doc, {"affaire": {}})
        self.assertEqual(len(doc.tables), 0)
        self.assertEqual(len(doc.sections), 1)


class EndpointTests(unittest.TestCase):
    def setUp(self) -> None:
        try:
            from fastapi.testclient import TestClient
        except Exception as e:  # fastapi ou httpx absent
            self.skipTest("fastapi/httpx indisponible : %s" % e)
        self.tmp = tempfile.TemporaryDirectory()
        os.environ["OUT_DIR"] = self.tmp.name
        os.environ["FICHES_DIR"] = self.tmp.name
        os.environ["LLM_SETTINGS_FILE"] = str(Path(self.tmp.name) / "llm_settings.json")
        sys.modules.pop("main", None)
        import main
        self.main = main
        self.client = TestClient(main.app)
        self.calls: list[tuple] = []

    def tearDown(self) -> None:
        self.tmp.cleanup()

    def _fake(self, reply: str):
        def fake(system, user, max_tokens=4000, provider="", model=""):
            self.calls.append((system, user, max_tokens))
            return reply
        self.main.llm_complete = fake

    def test_cctp_vide_400(self) -> None:
        r = self.client.post("/api/recap-cctp", json={"cctp_text": "  "})
        self.assertEqual(r.status_code, 400)

    def test_generation_ok(self) -> None:
        self._fake(json.dumps({"matrice": [LIGNE_MATRICE]}))
        r = self.client.post("/api/recap-cctp", json={"cctp_text": "Le titulaire fournit…",
                                                      "points": ["RPO"], "clarifications": []})
        self.assertEqual(r.status_code, 200)
        d = r.json()
        self.assertTrue(d["ok"])
        self.assertEqual(d["raw"], "")
        self.assertEqual(d["matrice"][0]["domaine"], "Sauvegarde")
        system, user, max_tokens = self.calls[0]
        self.assertTrue(system.startswith(self.main.ONEID_PREAMBLE))
        self.assertIn("- RPO", user)
        self.assertEqual(max_tokens, 6000)

    def test_reponse_inexploitable(self) -> None:
        self._fake("pas de JSON")
        d = self.client.post("/api/recap-cctp", json={"cctp_text": "x"}).json()
        self.assertFalse(d["ok"])
        self.assertEqual(d["raw"], "pas de JSON")


if __name__ == "__main__":
    unittest.main()

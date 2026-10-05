# /home/numa/projets/appel-offre/tests/test_recap_admin.py
"""Tests de l'étape « Focus administratif et contractuel ».

Couvre : normalisation de la sortie LLM (recap_admin), répartition du budget entre
documents, cohérence des colonnes backend / index.html, rendu Word (section paysage
partagée avec le récap technique) et endpoints /api/admin-docs/extract et
/api/recap-admin avec LLM simulé (ignorés si fastapi/httpx absents).

Exécution (depuis la racine du dépôt) : python -m unittest tests.test_recap_admin -v
"""
from __future__ import annotations

import json
import os
import re
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "app"))
sys.path.insert(0, str(ROOT / "app" / "engine"))

import recap_admin as ra  # noqa: E402

SORTIE = {
    "reference_marche": "M2026-30",
    "date_limite_questions": "05/10/2026 12h00",
    "checklist": [{"piece": "Acte d'engagement", "attendu_dce": "AE signé électroniquement",
                   "controle_oneid": "Vérifier la signature", "criticite": "élevée",
                   "responsable": "Direction", "echeance": "Avant dépôt", "statut": "a_preparer"}],
    "contractuel": [{"theme": "Pénalités", "clause": "Pénalités de retard plafonnées",
                     "impact_oneid": "Risque financier", "position": "Accepter", "niveau": "HIGH",
                     "source": "CCAP.pdf art. 11", "statut": "Validé"}],
    "questions": [{"num": "7", "question": "Pouvez-vous préciser la volumétrie ?", "enjeu": "Dimensionnement",
                   "priorite": "faible", "owner": "Avant-vente", "reponse_attendue": "Volumétrie en To",
                   "statut": "a_qualifier"},
                  {"num": "", "question": "Format du BPU ?", "statut": "a_qualifier"}],
}


class NormalizeAdminTests(unittest.TestCase):
    def test_parse_complet(self) -> None:
        r = ra.parse_admin("```json\n" + json.dumps(SORTIE, ensure_ascii=False) + "\n```")
        self.assertTrue(r["ok"])
        self.assertEqual(r["reference_marche"], "M2026-30")
        self.assertEqual(r["date_limite_questions"], "05/10/2026 12h00")
        self.assertEqual(r["date_limite_offres"], ra.A_PRECISER)
        self.assertEqual(r["checklist"][0]["criticite"], "Élevée")
        self.assertEqual(r["contractuel"][0]["niveau"], "Élevée")
        self.assertEqual(r["questions"][0]["priorite"], "Faible")

    def test_valide_interdit_au_llm_autorise_humain(self) -> None:
        self.assertEqual(ra.parse_admin(json.dumps(SORTIE))["contractuel"][0]["statut"], "a_confirmer")
        self.assertEqual(ra.normalize_admin(SORTIE, allow_valide=True)["contractuel"][0]["statut"], "valide")

    def test_questions_renumerotees_et_cellules_completees(self) -> None:
        q = ra.normalize_admin(SORTIE)["questions"]
        self.assertEqual([x["num"] for x in q], ["Q1", "Q2"])
        self.assertEqual(q[1]["enjeu"], ra.A_PRECISER)
        self.assertEqual(q[1]["priorite"], ra.NIVEAU_DEFAUT)

    def test_niveaux(self) -> None:
        for v, attendu in (("Haute", "Élevée"), ("moyen", "Moyenne"), ("basse", "Faible"), ("??", "À évaluer")):
            self.assertEqual(ra.normalize_niveau(v), attendu)

    def test_date_courte_pour_titre(self) -> None:
        self.assertEqual(ra.date_courte("05/10/2026 à 12h00 (plateforme PLACE, RC Art. 3)"), "05/10/2026 à 12h00")
        self.assertEqual(ra.date_courte("le 5/10/2026"), "5/10/2026")
        self.assertEqual(ra.date_courte("10 jours avant la remise"), "10 jours avant la remise")
        self.assertEqual(ra.date_courte(""), ra.A_PRECISER)

    def test_date_courte_meme_regle_en_js(self) -> None:
        html = (ROOT / "app" / "web" / "index.html").read_text(encoding="utf-8")
        js = re.search(r"function adminDateCourte\(v\)\{const m=String\(v\|\|''\)\.match\(/(.*?)/\);", html).group(1)
        self.assertEqual(js.replace("\\/", "/"), ra._DATE_RE.pattern)

    def test_sortie_inexploitable(self) -> None:
        r = ra.parse_admin("rien")
        self.assertFalse(r["ok"])
        self.assertEqual(r["reference_marche"], ra.A_PRECISER)

    def test_budget_reparti_ordre_conserve(self) -> None:
        docs = [("CCAP.pdf", "C" * 200000), ("RC.pdf", "R" * 1000), ("AE.docx", "A" * 50000)]
        m = ra.build_user_message(docs, cctp="T" * 20000)
        self.assertLess(m.index("CCAP.pdf"), m.index("RC.pdf"))
        self.assertLess(m.index("RC.pdf"), m.index("AE.docx"))
        self.assertIn("R" * 1000, m)                        # le document court passe en entier
        corps = re.split(r"===== [^\n]* =====\n", m)[1:4]   # textes des 3 documents (hors CCTP)
        self.assertEqual(sum(len(c.strip()) for c in corps), ra.MAX_INPUT)
        # RC (1 k) et AE (50 k) entiers ; le CCAP prend le reste du budget.
        self.assertEqual([len(c.strip()) for c in corps], [ra.MAX_INPUT - 51000, 1000, 50000])
        self.assertIn("T" * ra.MAX_CCTP_CONTEXT, m)
        self.assertNotIn("T" * (ra.MAX_CCTP_CONTEXT + 1), m)

    def test_budget_couvre_un_dce_courant(self) -> None:
        # Cas réel utilisateur 2026-10-05 : RC 27 k, CCAP 76 k, annexe 8 k, AE 8 k, annexe financière 62 k.
        tailles = [27324, 76121, 7764, 8312, 62113]
        m = ra.build_user_message([("D%d" % i, "x" * n) for i, n in enumerate(tailles)])
        self.assertEqual(m.count("x"), sum(tailles))         # rien n'est tronqué
        self.assertGreaterEqual(ra.MAX_INPUT, 240000)

    def test_documents_vides_ignores(self) -> None:
        self.assertNotIn("VIDE.pdf", ra.build_user_message([("VIDE.pdf", "  "), ("RC.pdf", "x")]))

    def test_prompt_conforme_au_contrat(self) -> None:
        for spec in ra.TABLES.values():
            for col, _ in spec["colonnes"]:
                self.assertIn('"%s"' % col, ra.RECAP_ADMIN_PROMPT)
        for k in ra.META_KEYS:
            self.assertIn('"%s"' % k, ra.RECAP_ADMIN_PROMPT)


class FrontendCoherenceTests(unittest.TestCase):
    def setUp(self) -> None:
        self.html = (ROOT / "app" / "web" / "index.html").read_text(encoding="utf-8")

    def test_etape_et_persistance(self) -> None:
        self.assertIn("const STEP_ORDER=['cctp','recap','admin',", self.html)
        self.assertIn("recap_admin:adminRec", self.html)
        self.assertIn("admin_docs:adminDocs", self.html)
        self.assertIn("'/api/admin-docs/extract'", self.html)
        self.assertIn("'/api/recap-admin'", self.html)
        self.assertIn('id="adminfile" accept=".pdf,.docx,.txt,.xlsx,.xlsm" multiple', self.html)

    def test_colonnes_identiques(self) -> None:
        block = self.html[self.html.index("const ADMIN_TABLES="):self.html.index("const ADMIN_WIDE=")]
        for key, spec in ra.TABLES.items():
            m = re.search(r"\{key:'%s',titre:'([^']*)',cols:\[(.*?)\]\}" % key, block)
            self.assertIsNotNone(m, key)
            self.assertEqual(m.group(1), spec["titre"])
            self.assertEqual(re.findall(r"\['(\w+)',", m.group(2)), [c for c, _ in spec["colonnes"]])
        self.assertNotIn("IMS", block)
        wide = re.search(r"const ADMIN_WIDE=\[(.*?)\];", self.html).group(1)
        self.assertEqual(tuple(re.findall(r"'(\w+)'", wide)), ra.WIDE_COLS)
        self.assertIn("let ADMIN_BUDGET=%d;" % ra.MAX_INPUT, self.html)
        self.assertIn('accept=".pdf,.docx,.txt,.xlsx,.xlsm"', self.html)
        niv = re.search(r"const ADMIN_NIVEAUX=\[(.*?)\];", self.html).group(1)
        self.assertEqual(tuple(re.findall(r"'([^']+)'", niv)), ra.NIVEAUX)


class WordExportTests(unittest.TestCase):
    def setUp(self) -> None:
        try:
            from docx import Document
            import generer_doc as gd
        except ImportError as e:
            self.skipTest("python-docx indisponible : %s" % e)
        self.Document, self.gd = Document, gd

    def test_section_admin(self) -> None:
        doc = self.Document()
        self.gd.section_recap_admin(doc, {"affaire": {"recap_admin": SORTIE}})
        self.gd._landscape_close(doc)
        self.assertEqual(len(doc.tables), 4)                 # bandeau Objectif + 3 tableaux
        self.assertEqual(doc.tables[0].rows[0].cells[0].text, ra.OBJECTIF)
        q = doc.tables[3]
        self.assertEqual(q.rows[0].cells[0].text, "Questions à déposer sur PLACE avant le 05/10/2026 12h00")
        self.assertEqual([c.text for c in q.rows[1].cells], [l for _, l in ra.TABLES["questions"]["colonnes"]])
        self.assertEqual(q.rows[2].cells[0].text, "Q1")
        self.assertIn("Focus administratif et contractuel – M2026-30", [p.text for p in doc.paragraphs])
        self.assertEqual(doc.tables[2].rows[2].cells[-1].text, "Validé")
        self.assertGreater(doc.sections[-2].page_width, doc.sections[-2].page_height)
        self.assertLess(doc.sections[-1].page_width, doc.sections[-1].page_height)

    def test_technique_puis_admin_une_seule_section_paysage(self) -> None:
        doc = self.Document()
        spec = {"affaire": {"recap_admin": SORTIE, "recap_cctp": {
            "matrice": [{"domaine": "Sauvegarde", "exigences": "x", "statut": "a_preparer"}]}}}
        self.gd.section_recap_cctp(doc, spec)
        self.gd.section_recap_admin(doc, spec)
        self.gd._landscape_close(doc)
        self.assertEqual(len(doc.sections), 3)               # portrait, paysage (les 2 focus), portrait

    def test_rien_si_vide(self) -> None:
        doc = self.Document()
        self.gd.section_recap_admin(doc, {"affaire": {"recap_admin": {"reference_marche": "X"}}})
        self.gd._landscape_close(doc)
        self.assertEqual((len(doc.tables), len(doc.sections)), (0, 1))


class EndpointTests(unittest.TestCase):
    def setUp(self) -> None:
        try:
            from fastapi.testclient import TestClient
        except Exception as e:
            self.skipTest("fastapi/httpx indisponible : %s" % e)
        self.tmp = tempfile.TemporaryDirectory()
        for k in ("OUT_DIR", "FICHES_DIR"):
            os.environ[k] = self.tmp.name
        os.environ["LLM_SETTINGS_FILE"] = str(Path(self.tmp.name) / "llm_settings.json")
        sys.modules.pop("main", None)
        import main
        self.main, self.client, self.calls = main, TestClient(main.app), []

    def tearDown(self) -> None:
        self.tmp.cleanup()

    def test_extraction_txt_et_doc_refuse(self) -> None:
        r = self.client.post("/api/admin-docs/extract",
                             files=[("files", ("RC.txt", "Marché M2026-30".encode(), "text/plain")),
                                    ("files", ("AE.txt", b"Acte", "text/plain"))])
        self.assertEqual(r.status_code, 200)
        self.assertEqual([d["name"] for d in r.json()["docs"]], ["RC.txt", "AE.txt"])
        self.assertEqual(r.json()["docs"][0]["text"], "Marché M2026-30")
        r = self.client.post("/api/admin-docs/extract", files=[("files", ("vieux.doc", b"x", "application/msword"))])
        self.assertEqual(r.status_code, 400)

    def test_extraction_xlsx_et_docx_avec_tableaux(self) -> None:
        import io
        import docx
        import openpyxl
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = "BPU"
        ws.append(["Poste", "Prix unitaire HT"]); ws.append(["Serveur", 1234.5]); ws.append([None, None])
        xb = io.BytesIO(); wb.save(xb)
        d = docx.Document(); d.add_paragraph("ACTE D'ENGAGEMENT")
        t = d.add_table(rows=1, cols=2); t.rows[0].cells[0].text = "Montant HT"; t.rows[0].cells[1].text = "à compléter"
        db = io.BytesIO(); d.save(db)
        r = self.client.post("/api/admin-docs/extract", files=[
            ("files", ("Annexe financière.xlsx", xb.getvalue(), "application/octet-stream")),
            ("files", ("AE.docx", db.getvalue(), "application/octet-stream"))])
        self.assertEqual(r.status_code, 200)
        docs = {x["name"]: x["text"] for x in r.json()["docs"]}
        self.assertEqual(docs["Annexe financière.xlsx"], "### Feuille : BPU\nPoste | Prix unitaire HT\nServeur | 1234.5")
        self.assertIn("Montant HT | à compléter", docs["AE.docx"])
        self.assertEqual(r.json()["budget"], ra.MAX_INPUT)
        for nom in ("vieux.xls", "image.png"):
            self.assertEqual(self.client.post("/api/admin-docs/extract",
                                              files=[("files", (nom, b"x", "application/octet-stream"))]).status_code, 400)

    def test_recap_admin(self) -> None:
        def fake(system, user, max_tokens=4000, provider="", model=""):
            self.calls.append((system, user, max_tokens))
            return json.dumps(SORTIE)
        self.main.llm_complete = fake
        self.assertEqual(self.client.post("/api/recap-admin", json={"docs": []}).status_code, 400)
        d = self.client.post("/api/recap-admin", json={"docs": [{"name": "RC.pdf", "text": "Règlement"}],
                                                       "cctp_text": "CCTP"}).json()
        self.assertTrue(d["ok"])
        self.assertEqual(d["reference_marche"], "M2026-30")
        system, user, mt = self.calls[0]
        self.assertTrue(system.startswith(self.main.ONEID_PREAMBLE))
        self.assertIn("===== DOCUMENT : RC.pdf =====", user)
        self.assertEqual(mt, 6000)


if __name__ == "__main__":
    unittest.main()
